import os
import psycopg2
import csv
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle
from reportlab.lib import colors
from conf_bd import host, user, password, db_name

from datetime import datetime


def db_upload():
    connection = psycopg2.connect(
        host=host,
        user=user,
        password=password,
        database=db_name
    )
    cursor = connection.cursor()

    tables = ["users", "alerts", "cryptocurrencies"]

    csv_folder = "csv"
    xlsx_folder = "xls"
    pdf_folder = "pdf"

    os.makedirs(csv_folder, exist_ok=True)
    os.makedirs(xlsx_folder, exist_ok=True)
    os.makedirs(pdf_folder, exist_ok=True)

    workbook = Workbook()
    xlsx_sheets = {}

    for table in tables:
        cursor.execute(f"SELECT * FROM {table};")
        data = cursor.fetchall()
        columns = [desc[0] for desc in cursor.description]

        csv_path = os.path.join(csv_folder, f"{table}.csv")
        with open(csv_path, "w", newline="", encoding="utf-8") as csv_file:
            writer = csv.writer(csv_file)
            writer.writerow(columns)
            writer.writerows(data)

        if not xlsx_sheets:
            ws = workbook.active
            ws.title = table
        else:
            ws = workbook.create_sheet(title=table)
        ws.append(columns)
        for row in data:
            ws.append(row)
        xlsx_sheets[table] = ws

    xlsx_path = os.path.join(xlsx_folder, "data.xlsx")
    workbook.save(xlsx_path)

    generate_pdf(cursor, pdf_folder)

    cursor.close()
    connection.close()

    print(f"Экспорт завершен! CSV сохранены в папке '{csv_folder}', XLSX в '{xlsx_folder}', PDF в '{pdf_folder}'.")


def generate_pdf(cursor, pdf_folder):
    # Запрашиваем у пользователя список криптовалют
    print("Введите названия криптовалют через запятую (например, BTC,ETH,USDT):")
    crypto_symbols_input = input()
    crypto_symbols = [symbol.strip() for symbol in crypto_symbols_input.split(",") if symbol.strip()]

    if not crypto_symbols:
        print("Ошибка: Вы не указали ни одной криптовалюты.")
        return

    # Формируем запрос с использованием списка криптовалют
    placeholders = ', '.join(['%s'] * len(crypto_symbols))  # Подготовка плейсхолдеров для списка
    query = f"""
        SELECT name, symbol, price, price_change_percentage_24h, volume_24h, last_updated
        FROM cryptocurrencies
        WHERE symbol IN ({placeholders});
    """
    params = crypto_symbols
    cursor.execute(query, params)
    filtered_data = cursor.fetchall()

    if not filtered_data:
        print("Нет данных, соответствующих заданным фильтрам.")
        return

    pdf_path = os.path.join(pdf_folder, "cryptocurrencies_report.pdf")
    pdf_canvas = canvas.Canvas(pdf_path, pagesize=A4)
    pdf_canvas.setTitle("Cryptocurrency Report")

    margin_x = 50
    margin_y = 800
    line_height = 20

    pdf_canvas.setFont("Helvetica-Bold", 16)
    pdf_canvas.drawString(margin_x, margin_y, "Cryptocurrency Report")
    pdf_canvas.setFont("Helvetica", 12)
    pdf_canvas.drawString(margin_x, margin_y - line_height, f"Filter: {', '.join(crypto_symbols)}")

    y_position = margin_y - 2 * line_height

    for crypto in filtered_data:
        if y_position < 100:
            pdf_canvas.showPage()
            y_position = margin_y
            pdf_canvas.setFont("Helvetica-Bold", 16)
            pdf_canvas.drawString(margin_x, margin_y, "Cryptocurrency Report")
            y_position -= 2 * line_height

        name, symbol, price, price_change, volume, last_updated = crypto

        pdf_canvas.setFont("Helvetica-Bold", 14)
        pdf_canvas.setFillColor(colors.darkblue)
        pdf_canvas.drawString(margin_x, y_position, f"{name} ({symbol})")
        y_position -= line_height

        pdf_canvas.setFont("Helvetica", 12)
        pdf_canvas.setFillColor(colors.black)
        pdf_canvas.drawString(margin_x, y_position, f"💰 Price: ${price:.2f}")
        y_position -= line_height

        pdf_canvas.drawString(margin_x, y_position, f"📈 Change in 24 hours: {price_change:.2f}%")
        y_position -= line_height

        pdf_canvas.drawString(margin_x, y_position, f"📊 Volume per 24h: ${volume:.2f}")
        y_position -= line_height

        pdf_canvas.drawString(margin_x, y_position, f"🕒 Updated: {last_updated}")
        y_position -= 2 * line_height

    pdf_canvas.save()
    print(f"PDF отчет сохранен: {pdf_path}")
